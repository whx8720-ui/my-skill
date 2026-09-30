#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
提示词巧匠 - Excel导出脚本

功能：
- 将提示词数据导出为Excel格式
- 支持详细版和简洁版两种格式
- 自动调整列宽和样式
- 支持单个提示词或批量导出

依赖：
- openpyxl (>=3.0.0)
"""

import json
import argparse
import sys
from datetime import datetime
from typing import List, Dict, Any
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误：缺少 openpyxl 依赖库")
    print("请执行：pip install openpyxl")
    sys.exit(1)


class ExcelExporter:
    """Excel导出器"""

    def __init__(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active

        # 定义样式
        self.header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        self.header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        self.data_font = Font(name='微软雅黑', size=10)
        self.data_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        self.border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

    def _apply_header_style(self, row: int):
        """应用表头样式"""
        for col in range(1, self.worksheet.max_column + 1):
            cell = self.worksheet.cell(row=row, column=col)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.border

    def _apply_data_style(self, start_row: int, end_row: int):
        """应用数据样式"""
        for row in range(start_row, end_row + 1):
            for col in range(1, self.worksheet.max_column + 1):
                cell = self.worksheet.cell(row=row, column=col)
                cell.font = self.data_font
                cell.alignment = self.data_alignment
                cell.border = self.border

    def _auto_adjust_column_width(self):
        """自动调整列宽"""
        for col in range(1, self.worksheet.max_column + 1):
            max_length = 0
            column_letter = get_column_letter(col)

            for row in range(1, self.worksheet.max_row + 1):
                cell_value = str(self.worksheet.cell(row=row, column=col).value or '')
                # 计算中文字符长度（每个中文字符算2个长度）
                length = 0
                for char in cell_value:
                    if ord(char) > 127:
                        length += 2
                    else:
                        length += 1
                max_length = max(max_length, length)

            # 设置列宽（最小10，最大60）
            adjusted_width = min(max(max_length + 2, 10), 60)
            self.worksheet.column_dimensions[column_letter].width = adjusted_width

    def _export_detailed_format(self, prompts: List[Dict[str, Any]]):
        """导出详细版（包含所有字段）"""
        # 表头
        headers = [
            '提示词ID', '使用场景', '目标导向', '效率优先', '准确性优先',
            '提示词全文', '评分', '生成模式', '创建时间'
        ]

        # 添加维度列（14个维度）
        dimension_columns = [
            '角色设定', '目标受众', '语言风格', '内容结构', '输出格式',
            '思维深度', '创意水平', '专业性', '实用性', '可操作性',
            '清晰度', '完整性', '灵活性', '适配性'
        ]

        all_headers = headers + dimension_columns

        # 写入表头
        for col, header in enumerate(all_headers, start=1):
            self.worksheet.cell(row=1, column=col, value=header)

        self._apply_header_style(1)

        # 写入数据
        row_num = 2
        for prompt in prompts:
            # 基本信息
            self.worksheet.cell(row=row_num, column=1, value=prompt.get('prompt_id', ''))
            self.worksheet.cell(row=row_num, column=2, value=prompt.get('scenario', ''))
            self.worksheet.cell(row=row_num, column=3, value=prompt.get('goal_orientation', ''))
            self.worksheet.cell(row=row_num, column=4, value='是' if prompt.get('is_efficiency_priority') else '否')
            self.worksheet.cell(row=row_num, column=5, value='是' if prompt.get('is_accuracy_priority') else '否')
            self.worksheet.cell(row=row_num, column=6, value=prompt.get('prompt_content', ''))
            self.worksheet.cell(row=row_num, column=7, value=prompt.get('score', ''))
            self.worksheet.cell(row=row_num, column=8, value=prompt.get('mode', ''))

            # 创建时间
            created_at = prompt.get('created_at', '')
            if created_at:
                try:
                    if isinstance(created_at, str):
                        # 尝试解析ISO格式时间
                        dt = datetime.fromisoformat(created_at.replace('Z', '+00:00'))
                        formatted_time = dt.strftime('%Y-%m-%d %H:%M:%S')
                    else:
                        formatted_time = str(created_at)
                except:
                    formatted_time = str(created_at)
            else:
                formatted_time = ''
            self.worksheet.cell(row=row_num, column=9, value=formatted_time)

            # 维度数据
            dimensions = prompt.get('dimensions', {})
            dimension_values = [
                dimensions.get('identity', ''),
                dimensions.get('target_audience', ''),
                dimensions.get('style', ''),
                dimensions.get('structure', ''),
                dimensions.get('output_format', ''),
                dimensions.get('thinking_depth', ''),
                dimensions.get('creativity', ''),
                dimensions.get('professionalism', ''),
                dimensions.get('practicality', ''),
                dimensions.get('operability', ''),
                dimensions.get('clarity', ''),
                dimensions.get('completeness', ''),
                dimensions.get('flexibility', ''),
                dimensions.get('adaptability', '')
            ]

            for col, value in enumerate(dimension_values, start=10):
                self.worksheet.cell(row=row_num, column=col, value=value)

            row_num += 1

        self._apply_data_style(2, row_num - 1)
        self._auto_adjust_column_width()

        # 设置行高
        self.worksheet.row_dimensions[1].height = 30
        for row in range(2, row_num):
            self.worksheet.row_dimensions[row].height = 60

    def _export_simple_format(self, prompts: List[Dict[str, Any]]):
        """导出简洁版（只包含关键字段）"""
        # 表头
        headers = [
            '提示词ID', '使用场景', '目标导向', '评分', '生成模式', '创建时间'
        ]

        # 写入表头
        for col, header in enumerate(headers, start=1):
            self.worksheet.cell(row=1, column=col, value=header)

        self._apply_header_style(1)

        # 写入数据
        row_num = 2
        for prompt in prompts:
            self.worksheet.cell(row=row_num, column=1, value=prompt.get('prompt_id', ''))
            self.worksheet.cell(row=row_num, column=2, value=prompt.get('scenario', ''))
            self.worksheet.cell(row=row_num, column=3, value=prompt.get('goal_orientation', ''))
            self.worksheet.cell(row=row_num, column=4, value=prompt.get('score', ''))
            self.worksheet.cell(row=row_num, column=5, value=prompt.get('mode', ''))

            # 创建时间
            created_at = prompt.get('created_at', '')
            if created_at:
                try:
                    if isinstance(created_at, str):
                        dt = datetime.fromisoformat(created_at.replace('Z', '+00:00'))
                        formatted_time = dt.strftime('%Y-%m-%d %H:%M:%S')
                    else:
                        formatted_time = str(created_at)
                except:
                    formatted_time = str(created_at)
            else:
                formatted_time = ''
            self.worksheet.cell(row=row_num, column=6, value=formatted_time)

            row_num += 1

        self._apply_data_style(2, row_num - 1)
        self._auto_adjust_column_width()

        # 设置行高
        self.worksheet.row_dimensions[1].height = 30
        for row in range(2, row_num):
            self.worksheet.row_dimensions[row].height = 25

    def export(self, prompts: List[Dict[str, Any]], output_path: str, format_type: str = 'detailed'):
        """
        导出提示词到Excel

        Args:
            prompts: 提示词数据列表
            output_path: 输出文件路径
            format_type: 格式类型（detailed详细版 / simple简洁版）
        """
        if not prompts:
            raise ValueError("提示词数据为空")

        # 根据格式类型导出
        if format_type == 'detailed':
            self._export_detailed_format(prompts)
            self.worksheet.title = "提示词详细记录"
        elif format_type == 'simple':
            self._export_simple_format(prompts)
            self.worksheet.title = "提示词简洁记录"
        else:
            raise ValueError(f"不支持的格式类型：{format_type}")

        # 保存文件
        self.workbook.save(output_path)
        print(f"✅ 导出成功：{output_path}")
        print(f"📊 共导出 {len(prompts)} 条记录")
        print(f"📄 格式类型：{'详细版' if format_type == 'detailed' else '简洁版'}")


def load_prompts_from_json(json_path: str) -> List[Dict[str, Any]]:
    """从JSON文件加载提示词数据"""
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        # 兼容多种数据格式
        if isinstance(data, list):
            return data
        elif isinstance(data, dict) and 'prompts' in data:
            return data['prompts']
        elif isinstance(data, dict):
            return [data]
        else:
            raise ValueError("不支持的数据格式")

    except FileNotFoundError:
        raise ValueError(f"文件不存在：{json_path}")
    except json.JSONDecodeError as e:
        raise ValueError(f"JSON解析错误：{str(e)}")


def main():
    """主函数"""
    parser = argparse.ArgumentParser(description='提示词巧匠 - Excel导出工具')
    parser.add_argument('--input', '-i', required=True, help='输入的JSON文件路径')
    parser.add_argument('--output', '-o', default='prompts_export.xlsx', help='输出的Excel文件路径（默认：prompts_export.xlsx）')
    parser.add_argument('--format', '-f', choices=['detailed', 'simple'], default='detailed',
                        help='导出格式：detailed详细版 / simple简洁版（默认：detailed）')

    args = parser.parse_args()

    # 加载数据
    print(f"📂 正在读取数据：{args.input}")
    prompts = load_prompts_from_json(args.input)

    # 导出Excel
    print(f"📊 正在导出Excel文件...")
    exporter = ExcelExporter()
    exporter.export(prompts, args.output, args.format)


if __name__ == '__main__':
    main()
